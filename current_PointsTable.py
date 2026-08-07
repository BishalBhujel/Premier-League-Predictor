# -------------------------------------------------------------------------------------------
# -----------------------------------Premier League Original Website-------------------------
# -------------------------------------------------------------------------------------------


# import requests
# from openpyxl import load_workbook, Workbook

# premier_league_standings_url = 'https://sdp-prem-prod.premier-league-prod.pulselive.com/api/v5/competitions/8/seasons/2026/standings?live=false'

# standings_table_data = requests.get(premier_league_standings_url)
# # wb = load_workbook('/Users/bishalbhujel/Downloads/team_data.xlsx')
# wb = Workbook('/Users/bishalbhujel/Downloads/team_data.xlsx')
# ws = wb.create_sheet('Current_Season_Standings')



# combined_key_last = []
# for data in standings_table_data.json()['tables'][0]['entries']:
#         combined_key = []
#         combined_data = []
#         for key, value in data.items():
#                 # print(value.keys())
#             # if key not in wb.sheetnames:
#                 # wb.create_sheet(key)
#             data_col = list(value.values())
#             keys_col = list(value.keys())
#             # ws = wb[key]
#             # ws.append(data_col)
#             for data in data_col:
#                 combined_data.append(data)
#             # print(keys_col)
#             if not combined_key_last:
#                 for key in keys_col:
#             #     if not combined_key:
#                     combined_key.append(key)
#         if not combined_key_last:
#             ws.append(combined_key)
#             combined_key_last = combined_key
#         ws.append(combined_data)
            

# wb.save('team_data.xlsx')



import requests
from openpyxl import load_workbook, Workbook
import os

# api point for the premier leaguge points table 
premier_league_standings_url = 'https://sdp-prem-prod.premier-league-prod.pulselive.com/api/v5/competitions/8/seasons/2026/standings?live=false'

standings_table_data = requests.get(premier_league_standings_url)
# creating a new excel workbook to store the table information
file_path = os.getcwd() + '/datasets/current_points_table.xlsx'
wb = Workbook(file_path)

combined_key = []
# creating a PK to uniquely identify records accross each sheet in the workbook
team_key = 1

for data in standings_table_data.json()['tables'][0]['entries']:
        combined_data = []
        for key, value in data.items():
            # Checking if the sheet already exists in the workbook, if not create a new sheet with the key name
            if key not in wb.sheetnames:
                wb.create_sheet(key)
            # Appending the data to the respective sheet in the workbook
            data_col = list(value.values()) + [team_key]
            keys_col = list(value.keys()) + ['team_key']
            ws = wb[key]
            # Appending the keys to the sheet only once, and then appending the data for each team
            if not combined_key:
                ws.append(keys_col)
            ws.append(data_col)
        # Incrementing the team_key for the next team record
        team_key += 1
        combined_key = ['a']
            
wb.save(file_path)