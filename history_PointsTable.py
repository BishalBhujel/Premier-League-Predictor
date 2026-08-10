import requests
import os
from openpyxl import Workbook

season_range = range(2008, 2026)

file_path = os.getcwd() + '/datasets/history_points_table.xlsx'
wb = Workbook(file_path)


for season in season_range:
    combined_key = []
    # creating a PK to uniquely identify records accross each sheet in the workbook
    team_key = 1
    url = f"https://sdp-prem-prod.premier-league-prod.pulselive.com/api/v5/competitions/8/seasons/{season}/standings?live=false"
    season_data = requests.get(url)

    for data in season_data.json()['tables'][0]['entries']:
        combined_data = []
        for key, value in data.items():
            # Checking if the sheet already exists in the workbook, if not create a new sheet with the key name
            if key+f"_{season}" not in wb.sheetnames:
                wb.create_sheet(key+f"_{season}")
            # Appending the data to the respective sheet in the workbook
            data_col = list(value.values()) + [team_key]
            keys_col = list(value.keys()) + ['team_key']
            ws = wb[key+f"_{season}"]
            # Appending the keys to the sheet only once, and then appending the data for each team
            if not combined_key:
                ws.append(keys_col)
            ws.append(data_col)
        # Incrementing the team_key for the next team record
        team_key += 1
        combined_key = ['a']
                
wb.save(file_path)